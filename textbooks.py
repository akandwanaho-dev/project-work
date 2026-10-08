"""Textbook parsing, indexing, and book management."""
import hashlib
import io
import re
from datetime import datetime
import streamlit as st
from docx import Document as DocxDocument
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from config import DEFAULT_EMBED_MODEL
from auth import save_book_to_db, delete_book_from_db

def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def safe_name(name: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._ -]", "_", name).strip()


def chapter_from(text, fallback="Unknown chapter"):
    for p in [
        r"(?im)^\s*(chapter\s+\d+(?:\.\d+)?\s*[:\-–]?\s*[^\n]{0,90})",
        r"(?im)^\s*(unit\s+\d+(?:\.\d+)?\s*[:\-–]?\s*[^\n]{0,90})",
        r"(?im)^\s*(topic\s+\d+(?:\.\d+)?\s*[:\-–]?\s*[^\n]{0,90})",
    ]:
        m = re.search(p, text[:4000])
        if m:
            return re.sub(r"\s+", " ", m.group(1)).strip()
    return fallback


def raw_docs_from_upload(upload):
    data = upload.getvalue()
    name = safe_name(upload.name)
    ext = name.rsplit(".", 1)[-1].lower()
    docs = []

    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        chapter = "Unknown chapter"
        for page_no, page in enumerate(reader.pages, start=1):
            text = clean(page.extract_text())
            if not text:
                continue
            chapter = chapter_from(text, chapter)
            docs.append(Document(page_content=text, metadata={"source": name, "page": page_no, "chapter": chapter, "type": ext}))

    elif ext == "docx":
        d = DocxDocument(io.BytesIO(data))
        chapter = "Unknown chapter"
        section = 1
        buf = []
        for para in d.paragraphs:
            t = clean(para.text)
            if not t:
                continue
            if re.match(r"(?i)^(chapter|unit|topic)\s+\d+", t):
                if buf:
                    docs.append(Document(page_content="\n".join(buf), metadata={"source": name, "page": section, "chapter": chapter, "type": ext}))
                    section += 1
                    buf = []
                chapter = t
            buf.append(t)
        if buf:
            docs.append(Document(page_content="\n".join(buf), metadata={"source": name, "page": section, "chapter": chapter, "type": ext}))

    elif ext in {"txt", "md"}:
        text = clean(data.decode("utf-8", errors="ignore"))
        parts = re.split(r"(?im)(?=^\s*(?:chapter|unit|topic)\s+\d+)", text)
        chapter = "Unknown chapter"
        section = 1
        for part in parts:
            part = clean(part)
            if not part:
                continue
            chapter = chapter_from(part, chapter)
            docs.append(Document(page_content=part, metadata={"source": name, "page": section, "chapter": chapter, "type": ext}))
            section += 1
    else:
        raise ValueError("Unsupported file type")

    return data, name, docs


def embeddings_model():
    return HuggingFaceEmbeddings(
        model_name=DEFAULT_EMBED_MODEL,
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True},
    )


def rebuild_index():
    if not st.session_state.documents:
        st.session_state.vectorstore = None
    else:
        st.session_state.vectorstore = FAISS.from_documents(st.session_state.documents, embeddings_model())


def index_uploads(files):
    existing = {m["hash"] for m in st.session_state.books.values()}
    added = 0
    for upload in files or []:
        data, name, raw_docs = raw_docs_from_upload(upload)
        h = sha(data)
        if h in existing:
            continue
        splitter = RecursiveCharacterTextSplitter(chunk_size=1100, chunk_overlap=180, separators=["\n\n", "\n", ". ", " ", ""])
        chunks = splitter.split_documents(raw_docs)
        for i, d in enumerate(chunks):
            d.metadata["hash"] = h
            d.metadata["chunk_id"] = f"{h[:12]}-{i}"
        st.session_state.documents.extend(chunks)
        st.session_state.books[name] = {
            "name": name,
            "hash": h,
            "size_mb": round(len(data)/(1024*1024), 2),
            "chunks": len(chunks),
            "pages": len({d.metadata.get("page") for d in raw_docs}),
            "chapters": sorted({d.metadata.get("chapter", "Unknown chapter") for d in chunks}),
            "uploaded": datetime.now().strftime("%Y-%m-%d %H:%M"),
        }
        existing.add(h)
        added += 1
    if added:
        rebuild_index()
    return added


def remove_book(name):
    st.session_state.documents = [d for d in st.session_state.documents if d.metadata.get("source") != name]
    st.session_state.books.pop(name, None)
    rebuild_index()

